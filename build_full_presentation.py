import copy
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Pt

FONT = "Noto Sans CJK KR"
C_TITLE = RGBColor(0x11, 0x11, 0x11)
C_SECBG = RGBColor(0xF2, 0xF2, 0xF2)
C_NUM = RGBColor(0x99, 0x99, 0x99)

EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
LINK = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link"


def copy_slide(dest_prs, src_slide):
    layout = dest_prs.slide_layouts[0]
    new_slide = dest_prs.slides.add_slide(layout)
    for shp in list(new_slide.shapes):
        shp._element.getparent().remove(shp._element)

    src_cSld = src_slide._element.find(qn("p:cSld"))
    bg = src_cSld.find(qn("p:bg"))
    new_cSld = new_slide._element.find(qn("p:cSld"))
    old_bg = new_cSld.find(qn("p:bg"))
    if old_bg is not None:
        new_cSld.remove(old_bg)
    if bg is not None:
        new_cSld.insert(0, copy.deepcopy(bg))

    rid_map = {}
    for rel_id, rel in src_slide.part.rels.items():
        if rel.reltype == RT.IMAGE:
            blob = rel.target_part.blob
            _, new_rid = new_slide.part.get_or_add_image_part(io.BytesIO(blob))
            rid_map[rel_id] = new_rid

    for el in list(src_slide.shapes._spTree):
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        new_el = copy.deepcopy(el)
        for node in new_el.iter():
            for attr_name in (EMBED, LINK):
                if attr_name in node.attrib:
                    old_rid = node.attrib[attr_name]
                    if old_rid in rid_map:
                        node.attrib[attr_name] = rid_map[old_rid]
        new_slide.shapes._spTree.append(new_el)
    return new_slide


def add_section_slide(prs, title, page_num, subtitle=None):
    layout = prs.slide_layouts[0]
    s = prs.slides.add_slide(layout)
    for shp in list(s.shapes):
        shp._element.getparent().remove(shp._element)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C_SECBG

    tb = s.shapes.add_textbox(Emu(548640), Emu(1828800), Emu(8046720), Emu(1371600))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.name = FONT
    run.font.color.rgb = C_TITLE

    if subtitle:
        tb2 = s.shapes.add_textbox(Emu(548640), Emu(3111120), Emu(8046720), Emu(457200))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(18)
        run2.font.name = FONT
        run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    add_page_num(s, page_num)
    return s


def add_page_num(slide, n):
    tb = slide.shapes.add_textbox(Emu(8503920), Emu(4754880), Emu(457200), Emu(274320))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = str(n)
    run.font.size = Pt(12)
    run.font.name = FONT
    run.font.color.rgb = C_NUM


def renumber_page(slide, n):
    # page-number textbox is the bottom-right small numeric text added by every build script
    for shp in slide.shapes:
        if not shp.has_text_frame:
            continue
        if shp.left == 8503920 and shp.top == 4754880:
            tf = shp.text_frame
            for p in tf.paragraphs:
                for r in p.runs:
                    if r.text.strip().isdigit():
                        r.text = str(n)
            return True
    return False


prs = Presentation("section1_intro.pptx")
page = 0
for slide in prs.slides:
    page += 1
    renumber_page(slide, page)

for fname in ["section2_tdma_protocol.pptx", "section3_block_diagrams.pptx"]:
    src = Presentation(fname)
    for slide in src.slides:
        new_slide = copy_slide(prs, slide)
        page += 1
        renumber_page(new_slide, page)

page += 1
add_section_slide(prs, "4. TDMA 통신 시연", page)
page += 1
add_section_slide(prs, "5. 마무리", page)

prs.save("presentation_full.pptx")
print("done:", page, "slides")
